"""Import candidate evidence from the Fintech source workbook."""

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from clearpass.contracts.models import EvidenceRecord

SOURCE = Path("data/evidence/source/evidence.xlsx")
DESTINATION = Path("data/evidence/evidence.json")

HEADERS = (
    "evidence_id",
    "atomic_proposition",
    "verbatim_supporting_passage",
    "source_name",
    "source_type",
    "source_url",
    "passage_location",
    "retrieved_on",
    "scope",
    "limitations",
    "review_status",
    "reviewed_by",
    "priority",
    "source_document_version",
    "retrieval_tags",
    "follow_up_question",
    "human_reviewer",
    "human_reviewed_on",
)


def text_or_none(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def date_or_none(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return text_or_none(value)


def import_workbook(source=SOURCE):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Evidence"]
        rows = sheet.iter_rows(values_only=True)

        for row in rows:
            if tuple(row[:18]) == HEADERS:
                break
        else:
            raise ValueError("Evidence header row not found")

        results = []
        seen_ids = set()

        for row_number, row in enumerate(rows, start=6):
            if not any(value is not None for value in row):
                continue

            if not row[0]:
                raise ValueError(f"Row {row_number}: missing evidence ID")

            data = dict(zip(HEADERS, row, strict=True))
            evidence_id = text_or_none(data["evidence_id"])

            if evidence_id in seen_ids:
                raise ValueError(f"Duplicate evidence ID: {evidence_id}")

            seen_ids.add(evidence_id)

            core = {key: data[key] for key in EvidenceRecord.model_fields}
            core["retrieved_on"] = date_or_none(core["retrieved_on"])

            validated = EvidenceRecord.model_validate(core)

            review_status = text_or_none(data["review_status"])
            human_reviewer = text_or_none(data["human_reviewer"])
            human_reviewed_on = date_or_none(data["human_reviewed_on"])

            # A status label alone is not proof of review.
            approved = (
                review_status == "human_reviewed_approved"
                and bool(human_reviewer)
                and bool(human_reviewed_on)
            )

            record = validated.model_dump(mode="json")
            record.update(
                {
                    "source_type": text_or_none(data["source_type"]),
                    "review_status": review_status,
                    "human_reviewer": human_reviewer,
                    "human_reviewed_on": human_reviewed_on,
                    "approved_for_verification": approved,
                    "retrieval_tags": text_or_none(data["retrieval_tags"]),
                }
            )
            results.append(record)

        return results
    finally:
        workbook.close()


def main():
    records = import_workbook()

    if not records:
        raise ValueError("No evidence records found")

    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    DESTINATION.write_text(
        json.dumps(records, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    approved = sum(r["approved_for_verification"] for r in records)
    print(f"Imported {len(records)} structurally valid records")
    print(f"Approved for verification: {approved}")
    print(f"Awaiting review: {len(records) - approved}")
    print(f"Output: {DESTINATION}")


if __name__ == "__main__":
    main()
